#!/usr/bin/env python3
"""
Generates the authoritative FWA Metric to Source Mapping spreadsheet for the 2024 POPCEN & PSA FIES Edition.
Saved directly into FWA Design Phase 2024 POPCEN/.
"""

import os
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT_DIR = "FWA Design Phase 2024 POPCEN"

def build_mapping_workbook_2024():
    wb = openpyxl.Workbook()
    
    navy_dark = "1B365D"
    gray_border = "CBD5E1"
    
    fill_header = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    font_title = Font(name="Calibri", size=16, bold=True, color=navy_dark)
    font_sub = Font(name="Calibri", size=10, italic=True, color="64748B")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_category = Font(name="Calibri", size=10, bold=True, color=navy_dark)
    font_body = Font(name="Calibri", size=9.5)
    font_bold = Font(name="Calibri", size=9.5, bold=True)
    font_code = Font(name="Consolas", size=9, color="0F172A")
    
    border_cell = Border(
        left=Side(style="thin", color=gray_border),
        right=Side(style="thin", color=gray_border),
        top=Side(style="thin", color=gray_border),
        bottom=Side(style="thin", color=gray_border)
    )
    border_header = Border(
        left=Side(style="thin", color="FFFFFF"),
        right=Side(style="thin", color="FFFFFF"),
        top=Side(style="medium", color=navy_dark),
        bottom=Side(style="medium", color=navy_dark)
    )
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # TAB 1: Metric to Source Mapping
    ws1 = wb.active
    ws1.title = "Metric to Source Mapping (2024)"
    ws1.views.sheetView[0].showGridLines = True
    
    ws1["A1"] = "Philippine 5G FWA Rollout Model — Metric to Data Source Mapping (2024 POPCEN & FIES Edition)"
    ws1["A1"].font = font_title
    ws1["A2"] = "Complete data lineage citing the 2024 Census of Population (Proclamation No. 973), PSA FIES Poverty Statistics, and Converge Backbone GIS."
    ws1["A2"].font = font_sub
    
    headers_t1 = [
        "Metric / Column Header",
        "Category",
        "Authoritative Agency / Operator",
        "Primary Source Document / Asset",
        "Technical Derivation & Mathematical Formula",
        "Operational & Business Rationale"
    ]
    
    for col_idx, h in enumerate(headers_t1, 1):
        cell = ws1.cell(4, col_idx, h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_header
    ws1.row_dimensions[4].height = 28
    
    metrics_data = [
        (
            "PSGC Code",
            "Demographic & Admin",
            "Philippine Statistics Authority (PSA)",
            "PSA List of Barangays & PSGC Registry",
            "10-digit standardized geographic code (e.g., '0102801001') assigned by the Philippine Standard Geographic Code registry.",
            "Serves as the unique primary key joining census data, DICT coordinates, and municipal revenue class records."
        ),
        (
            "Region, Province, Municipality, Barangay",
            "Demographic & Admin",
            "Philippine Statistics Authority (PSA)",
            "PSA Official Administrative Hierarchy (DILG/PSA)",
            "Official administrative naming hierarchy as defined by the Republic of the Philippines DILG and PSA.",
            "Establishes precise administrative boundaries, political jurisdictions, and regional sales distribution territories."
        ),
        (
            "2024 Pop (Population)",
            "Demographic & Market Size",
            "Philippine Statistics Authority (PSA)",
            "2024 Census of Population (POPCEN), Proclamation No. 973 & Statistical Table.xlsx",
            "Official 2024 POPCEN population counts declared official by President Ferdinand R. Marcos Jr. (112,729,484 national total).",
            "Direct measure of market size and local human density; foundational input for calculating total addressable home broadband households."
        ),
        (
            "Urban / Rural ('U' / 'R')",
            "Demographic & Siting",
            "Philippine Statistics Authority (PSA)",
            "PSA Urban-Rural Classification of Barangays (CPH/POPCEN)",
            "Binary statistical classification based on population threshold, contiguous street pattern, and commercial establishments.",
            "Rural catchments ('R') receive priority under the Necessity pillar to target unserved populations with zero fixed-line broadband."
        ),
        (
            "Avg HH Size (Persons/HH)",
            "Demographic & Household",
            "PSA 2024 POPCEN (Statistical Table.xlsx, Table 1)",
            "2024 POPCEN Provincial Average Household Size Matrix",
            "Localized average persons per household derived from official 2024 POPCEN tables. National mean is 3.80 (updated from 4.10 in 2020).",
            "Converts raw population into family dwelling units (households). Using 2024 localized multipliers reflects smaller modern family sizes."
        ),
        (
            "Households",
            "Market Sizing",
            "Derived Metric (2024 POPCEN)",
            "Modeled: = ROUND(2024 Pop / Avg HH Size, 0)",
            "Mathematical division of barangay population by 2024 average household size, rounded to the nearest whole household.",
            "Represents the Total Addressable Market (TAM) of home broadband premises within the barangay boundaries."
        ),
        (
            "FIES Poverty Incidence (%)",
            "Economic & Poverty Metric",
            "Philippine Statistics Authority (PSA)",
            "PSA Family Income and Expenditure Survey (FIES) Official Poverty Statistics",
            "Proportion of families with per-capita income below the official poverty threshold (₱13,873/mo for family of 5; ranges from 1.1% in NCR to 23.5% in BARMM).",
            "Dual role: in Necessity, prioritizes unserved lower-income catchments for universal service; in Viability, discounts subsistence households to reflect commercial affordability."
        ),
        (
            "Municipal Income Class",
            "Economic & Policy",
            "Department of Finance (DOF) / BLGF",
            "DOF Department Order No. 074-2024 & RA 11964 (LGU Income Classification Act)",
            "Official classification of local government units from 1st class (highest income) to 6th class (lowest income) based on regular revenues.",
            "Focuses universal service on 4th to 6th class municipalities where wireline FTTH economics fail for incumbents."
        ),
        (
            "Latitude & Longitude",
            "Geospatial Coordinates",
            "Department of Information and Communications Technology (DICT)",
            "2026 DICT FPIAP GIDA Barangay Prioritization Tool Table.csv",
            "WGS84 decimal degree coordinates for 41,981 unique barangays, verified through DICT Free Public Internet Access Program (FPIAP) surveys.",
            "Provides the precise spatial centroid of the barangay settlement core for point-to-point distance calculations to fiber nodes and towers."
        ),
        (
            "Dist. to Nearest Node (km)",
            "Optical Transport Proximity",
            "Converge ICT Solutions Inc.",
            "NATIONAL & REGIONAL BACKBONE_NOV 2024_REPORT.kmz (2,405 active backbone / IP-MPLS nodes)",
            "Great-circle distance calculated via Haversine formula:\nd = 2R * arcsin(sqrt(sin^2(dphi/2) + cos(p1)*cos(p2)*sin^2(dlmb/2)))\nagainst all 2,405 Converge backbone nodes.",
            "Primary driver of backhaul CAPEX and deployment velocity. Sites within 1 km (Tier 1A) allow rapid fiber drop with minimal lateral construction cost."
        ),
        (
            "Dist. to Fiber Line (km)",
            "Optical Corridor Proximity",
            "Converge ICT Solutions Inc.",
            "NATIONAL & REGIONAL BACKBONE_NOV 2024_REPORT.kmz (4,401 fiber segments / 334,462 vertices)",
            "Geometric distance calculated via 0.25-degree spatial grid index against 44,959 downsampled fiber route vertices.",
            "Identifies barangays situated along existing long-haul optical transport corridors; awards +10 bonus points if within 2.0 km."
        ),
        (
            "Deploy Tier",
            "Engineering Readiness",
            "TelTech & Comclark Engineering",
            "Engineering Siting Specifications & Distance Matrix",
            "Categorical classification derived from dist_node_km and dist_line_km: Tier 1A (<1 km), Tier 1B (1-3 km), Tier 2 (3-5 km or <=2km line), Tier 3 (5-10 km), Tier 4 (>10 km)",
            "Defines civil works readiness, permitting timeline, and backhaul transmission architecture (direct fiber vs. microwave/satellite spur)."
        ),
        (
            "Carrier Spectrum",
            "Radio Spectrum",
            "3GPP & National Telecommunications Commission (NTC)",
            "3GPP TS 38.101-1 (Band n50, 1432-1517 MHz, 1.5 GHz L-Band TDD) & NTC Spectrum Allocations",
            "60 MHz contiguous TDD carrier bandwidth in sub-2 GHz L-Band spectrum.",
            "Provides the optimal compromise between deep indoor building penetration, non-line-of-sight (NLOS) propagation, and 2.5-3.5 km rural cell coverage radius."
        ),
        (
            "Subs @ 30% (Target Subscribers)",
            "Commercial Demand",
            "Commercial Business Model Assumption",
            "Commercial Business Case Assumption (FWA Meeting Pack)",
            "Modeled: = ROUND(Households * 0.30, 0)",
            "Calculates the commercial target subscriber base assuming a 30% steady-state adoption ceiling among addressable households. Defines subscriber loading requirements for dimensioning RAN capacity."
        ),
        (
            "Initial Day 1 Build (BTS)",
            "CAPEX & Siting Policy",
            "TelTech / Comclark Commercial Plan",
            "Option A Architecture Specification (FWA Design Phase)",
            "Fixed standard: Exactly 1 BTS installed per priority barangay on Day 1.",
            "Establishes immediate commercial presence, fast time-to-market, and coverage across 1,000 distinct communities per phase without speculative overbuilding."
        ),
        (
            "Peak 30% Demand Capacity (BTS)",
            "RAN Capacity Dimensioning",
            "3GPP / TelTech RAN Engineering",
            "Band n50 60MHz Dimensioning Model (@ 1,000 Subs/BTS)",
            "Modeled: = ROUNDUP(Target Subs @ 30% / 1,000, 0)\nDimensioned strictly for 1,000 subscribers per BTS total across 3 sectors (~333 subs/sector).",
            "Defines the long-term radio absorption capacity needed if 30% penetration is reached. Unlocks organic 2nd/3rd BTS on existing tower footprint without new land acquisition."
        ),
        (
            "Deploy Score (35%)",
            "Site Siting Pillar 3",
            "Multi-Criteria Spatial Engine",
            "Converge KMZ Distance Model & Haversine Matrix",
            "Exponential decay model:\nScore = 100 * exp(-dist_node_km / 3.0) + [10 pt bonus if dist_line_km <= 2.0 km].\nBounded [15.0, 100.0].",
            "Quantifies the ease and cost of site construction, presence of 220V grid power, existing tower infrastructure, and short-distance optical fiber backhaul."
        ),
        (
            "Viability Score (35%)",
            "Site Siting Pillar 2",
            "Commercial Financial Engine",
            "Household Demand Density & FIES Affordability Adjustment",
            "Continuous sweet-spot curve on poverty-adjusted subscribers:\nviable_subs = subs_target * (1 - 0.40 * poverty_pct/100)\nEvaluates 300 to 1,000 subs sweet-spot fit, adjusted for regional GDP.",
            "Guarantees that 1 BTS generates sufficient cash flow to achieve full CAPEX payback within 14 to 22 months while adjusting for local household purchasing power."
        ),
        (
            "Necessity Score (30%)",
            "Site Siting Pillar 1",
            "Broadband Deficit & Universal Service Engine",
            "PSA Rural Classification, DOF Income Class, Node Proximity & FIES Table A",
            "Multi-factor broadband deficit:\nScore = 0.40 * Rural_Comp + 0.25 * Class_Comp + 0.20 * OffNet_Comp + 0.15 * Poverty_Comp\nWhere OffNet_Comp is spatial distance (25 to 100 pts), and Poverty_Comp is scaled FIES poverty rate.",
            "Directly identifies communities suffering from lack of commercial broadband and economic digital exclusion, prioritizing unserved rural areas."
        ),
        (
            "Composite Score (0 - 100)",
            "Master Ranking",
            "Multi-Criteria Decision Analysis (MCDA)",
            "Master Prioritization Algorithm (generate_barangay_rollout_2024.py)",
            "Weighted Multi-Criteria Siting Algorithm:\nScore = 0.35 * Deploy_Score + 0.35 * Viability_Score + 0.30 * Necessity_Score",
            "The master sorting index ranking all 41,976 barangays nationwide into sequential 1,000-site commercial rollout batches."
        )
    ]
    
    for row_idx, item in enumerate(metrics_data, 5):
        fill_curr = fill_zebra if row_idx % 2 == 1 else PatternFill(fill_type=None)
        
        c1 = ws1.cell(row_idx, 1, item[0]); c1.font = font_bold; c1.alignment = align_left; c1.border = border_cell; c1.fill = fill_curr
        c2 = ws1.cell(row_idx, 2, item[1]); c2.font = font_category; c2.alignment = align_center; c2.border = border_cell; c2.fill = fill_curr
        c3 = ws1.cell(row_idx, 3, item[2]); c3.font = font_body; c3.alignment = align_left; c3.border = border_cell; c3.fill = fill_curr
        c4 = ws1.cell(row_idx, 4, item[3]); c4.font = font_code; c4.alignment = align_left; c4.border = border_cell; c4.fill = fill_curr
        c5 = ws1.cell(row_idx, 5, item[4]); c5.font = font_body; c5.alignment = align_left; c5.border = border_cell; c5.fill = fill_curr
        c6 = ws1.cell(row_idx, 6, item[5]); c6.font = font_body; c6.alignment = align_left; c6.border = border_cell; c6.fill = fill_curr
        ws1.row_dimensions[row_idx].height = 42
        
    ws1.column_dimensions["A"].width = 24
    ws1.column_dimensions["B"].width = 22
    ws1.column_dimensions["C"].width = 28
    ws1.column_dimensions["D"].width = 34
    ws1.column_dimensions["E"].width = 46
    ws1.column_dimensions["F"].width = 46
    ws1.freeze_panes = "A5"

    # TAB 2: Source Datasets
    ws2 = wb.create_sheet(title="Source Datasets (2024)")
    ws2.views.sheetView[0].showGridLines = True
    ws2["A1"] = "Primary Datasets Ingested (2024 POPCEN & PSA FIES)"
    ws2["A1"].font = font_title
    ws2["A2"] = "Complete inventory of raw datasets, census tables, and municipal records in FWA Design Phase 2024 POPCEN/data/."
    ws2["A2"].font = font_sub
    
    headers_t2 = [
        "Dataset Asset Name", "File Format & Size", "Authoring Agency / Entity",
        "Record Scope & Volume", "Key Fields Ingested", "Primary Operational Role in Model"
    ]
    for col_idx, h in enumerate(headers_t2, 1):
        cell = ws2.cell(4, col_idx, h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_header
    ws2.row_dimensions[4].height = 28
    
    datasets_data = [
        (
            "Statistical Table.xlsx (Table 1)",
            "Excel Workbook (.xlsx)\n20 KB",
            "Philippine Statistics Authority (PSA)\nNational Government",
            "118 Provinces & Highly Urbanized Cities\n(112,729,484 National Population)",
            "Total Population, Household Population, Number of Households, Average Household Size (2024)",
            "Official 2024 POPCEN counts declared by President Marcos Jr. via Proclamation No. 973. Updates population and modern household size (3.80)."
        ),
        (
            "NCR_Statistical Table_R13.xlsx",
            "Excel Workbook (.xlsx)\n113 KB",
            "Philippine Statistics Authority (PSA)",
            "1,728 barangays across 16 cities & 1 municipality in Metro Manila",
            "Barangay name, Total Population, Household Population, Number of Households (01 July 2024)",
            "Direct ground-truth 2024 barangay counts for the National Capital Region."
        ),
        (
            "Table A.xlsx & Table C.xlsx (FIES)",
            "Excel Workbooks (.xlsx)\n32 KB",
            "Philippine Statistics Authority (PSA)",
            "National income deciles & poverty statistics across 29.56M families",
            "Income Deciles, Number of Families, Average Annual Family Income & Expenditure, Poverty Threshold",
            "Provides empirical foundation for commercial affordability adjustments and universal service necessity prioritization."
        ),
        (
            "PSA List of Barangays.xlsx",
            "Excel Workbook (.xlsx)\n4.8 MB",
            "Philippine Statistics Authority (PSA)",
            "41,976 valid barangays\n(100% of Philippine territory)",
            "New 10 digit PSGC, Region, Province, Municipality, Barangay, Urban/Rural classification",
            "Master administrative backbone and geographic code structure joining all Philippine barangays."
        ),
        (
            "NATIONAL & REGIONAL BACKBONE_NOV 2024_REPORT.kmz",
            "Google Earth KMZ (.kmz)\n4.4 MB (48 MB uncompressed)",
            "Converge ICT Solutions Inc.",
            "2,405 active backbone nodes\n4,401 fiber line segments / 334,462 vertices",
            "Node coordinates (<Point>), Site Names, Fiber line routes (<LineString>)",
            "Ground-truth optical backbone transport map. Used to derive exact point-to-point fiber distances."
        ),
        (
            "2026 DICT FPIAP GIDA Barangay Prioritization Tool Table.csv",
            "Comma-Separated Values (.csv)\n5.2 MB",
            "Department of Information and Communications Technology (DICT)",
            "41,981 unique barangay coordinates nationwide",
            "Province, Locality, Barangay, Latitude, Longitude, GIDA Score",
            "Verified geospatial coordinates (WGS84) for local barangay centers, anchoring spatial distance calculations."
        )
    ]
    
    for row_idx, item in enumerate(datasets_data, 5):
        fill_curr = fill_zebra if row_idx % 2 == 1 else PatternFill(fill_type=None)
        c1 = ws2.cell(row_idx, 1, item[0]); c1.font = font_bold; c1.alignment = align_left; c1.border = border_cell; c1.fill = fill_curr
        c2 = ws2.cell(row_idx, 2, item[1]); c2.font = font_code; c2.alignment = align_center; c2.border = border_cell; c2.fill = fill_curr
        c3 = ws2.cell(row_idx, 3, item[2]); c3.font = font_body; c3.alignment = align_left; c3.border = border_cell; c3.fill = fill_curr
        c4 = ws2.cell(row_idx, 4, item[3]); c4.font = font_body; c4.alignment = align_left; c4.border = border_cell; c4.fill = fill_curr
        c5 = ws2.cell(row_idx, 5, item[4]); c5.font = font_body; c5.alignment = align_left; c5.border = border_cell; c5.fill = fill_curr
        c6 = ws2.cell(row_idx, 6, item[5]); c6.font = font_body; c6.alignment = align_left; c6.border = border_cell; c6.fill = fill_curr
        ws2.row_dimensions[row_idx].height = 48
        
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 22
    ws2.column_dimensions["C"].width = 26
    ws2.column_dimensions["D"].width = 28
    ws2.column_dimensions["E"].width = 38
    ws2.column_dimensions["F"].width = 46
    ws2.freeze_panes = "A5"

    xlsx_path = os.path.join(OUT_DIR, "FWA_Metric_to_Source_Mapping_2024.xlsx")
    wb.save(xlsx_path)
    print(f"  -> Saved {xlsx_path} ({os.path.getsize(xlsx_path):,} bytes).")
    
    csv_path = os.path.join(OUT_DIR, "FWA_Metric_to_Source_Mapping_2024.csv")
    with open(csv_path, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers_t1)
        for item in metrics_data:
            writer.writerow(item)
    print(f"  -> Exported CSV: {csv_path} ({os.path.getsize(csv_path):,} bytes).")

if __name__ == '__main__':
    build_mapping_workbook_2024()
